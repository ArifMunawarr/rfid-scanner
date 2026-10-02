import argparse
import os
import re
import signal
import socket
import sys
from datetime import datetime

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

try:
    from dotenv import load_dotenv
    load_dotenv()
except ImportError:
    pass

START_COMMAND = bytes.fromhex("BB 17 02 00 00 19 0D 0A")
STOP_COMMAND = bytes.fromhex("BB 18 00 18 0D 0A")

# Dict untuk track tag_id -> data tag yang disimpan pada sheet aktif.
seen_tags = {}
existing_tags = set()
session_new_tags = set()
TAGS_FILE = os.getenv("RFID_TAGS_FILE", "tags.xlsx")
TAG_MIN = os.getenv("RFID_TAG_MIN", "")
TAG_MAX = os.getenv("RFID_TAG_MAX", "")
CURRENT_SHEET_NAME = None
COMPARISON_SHEET_NAMES = []
NO_SAVE = False

# Tag ID yang dikecualikan (tidak diproses / diabaikan)
EXCLUDED_TAGS = {"0525111566", "0525111567", "0525111556"}

HEADERS = ["No", "Tag ID", "Device ID"]

def normalize_tag_id(value):
    """Normalisasi Tag ID Excel ke format teks 10 karakter."""
    if value is None:
        return None

    if isinstance(value, float) and value.is_integer():
        value = int(value)

    if isinstance(value, int):
        return str(value).zfill(10)

    value = str(value).strip()
    return value or None


def read_tags_from_sheet(ws):
    """Ambil seluruh Tag ID unik dari sebuah sheet."""
    tags = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        tag_id = normalize_tag_id(row[1] if len(row) > 1 else None)
        if tag_id:
            tags.add(tag_id)
    return tags


def read_tag_records_from_sheet(ws):
    """Ambil data Tag ID dari sheet agar mode lanjutan tidak menghapus data lama."""
    records = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        tag_id = normalize_tag_id(row[1] if len(row) > 1 else None)
        if not tag_id or tag_id in records:
            continue

        records[tag_id] = {
            "device_id": row[2] if len(row) > 2 else None,
        }
    return records


def make_unique_sheet_name(wb, requested_name=None):
    """Buat nama sheet valid dan unik untuk sesi scan saat ini."""
    base_name = requested_name or datetime.now().strftime("Scan %Y%m%d %H%M%S")
    base_name = re.sub(r"[\\/*?:\[\]]", "-", base_name).strip(" '") or "Scan"
    base_name = base_name[:31]
    existing_names = {name.lower() for name in wb.sheetnames}

    candidate = base_name
    suffix = 2
    while candidate.lower() in existing_names:
        suffix_text = f" ({suffix})"
        candidate = f"{base_name[:31 - len(suffix_text)]}{suffix_text}"
        suffix += 1
    return candidate


def format_sheet(ws):
    """Terapkan header dan ukuran kolom pada sheet scan."""
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_align = Alignment(horizontal="center")

    for col, title in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 22


def remove_rssi_columns(wb):
    """Hapus kolom berjudul RSSI dari seluruh sheet pada workbook."""
    changed_sheets = []

    for ws in wb.worksheets:
        rssi_columns = [
            cell.column
            for cell in ws[1]
            if str(cell.value or "").strip().upper() == "RSSI"
        ]
        if not rssi_columns:
            continue

        for column_index in reversed(rssi_columns):
            ws.delete_cols(column_index, 1)
        format_sheet(ws)
        changed_sheets.append(ws.title)

    return changed_sheets


def remove_rssi_from_existing_workbook():
    """Migrasikan workbook lama tanpa menjalankan proses scan."""
    if not os.path.exists(TAGS_FILE):
        raise FileNotFoundError(f"File {os.path.abspath(TAGS_FILE)} belum ada")

    wb = openpyxl.load_workbook(TAGS_FILE)
    try:
        changed_sheets = remove_rssi_columns(wb)
        if changed_sheets:
            wb.save(TAGS_FILE)
    finally:
        wb.close()

    return changed_sheets


def initialize_scan_sheet(requested_name=None, continue_sheet=None):
    """Siapkan sheet baru atau lanjutkan sheet lama tanpa menghapus datanya."""
    global CURRENT_SHEET_NAME, COMPARISON_SHEET_NAMES, existing_tags

    seen_tags.clear()
    session_new_tags.clear()

    if os.path.exists(TAGS_FILE):
        wb = openpyxl.load_workbook(TAGS_FILE)
        changed_sheets = remove_rssi_columns(wb)
        if changed_sheets:
            wb.save(TAGS_FILE)
            log(
                "Kolom RSSI dihapus dari sheet lama: "
                + ", ".join(changed_sheets)
            )
        COMPARISON_SHEET_NAMES = wb.sheetnames[:]
        existing_tags = set()
        for existing_ws in wb.worksheets:
            existing_tags.update(read_tags_from_sheet(existing_ws))

        if continue_sheet:
            if continue_sheet not in wb.sheetnames:
                available = ", ".join(wb.sheetnames)
                wb.close()
                raise ValueError(
                    f"Sheet '{continue_sheet}' tidak ditemukan. "
                    f"Sheet tersedia: {available}"
                )

            CURRENT_SHEET_NAME = continue_sheet
            ws = wb[CURRENT_SHEET_NAME]
            seen_tags.update(read_tag_records_from_sheet(ws))
        else:
            CURRENT_SHEET_NAME = make_unique_sheet_name(wb, requested_name)
            ws = wb.create_sheet(CURRENT_SHEET_NAME)
            format_sheet(ws)
            wb.save(TAGS_FILE)
    else:
        if continue_sheet:
            raise ValueError(
                f"Tidak dapat melanjutkan sheet '{continue_sheet}': "
                f"file {os.path.abspath(TAGS_FILE)} belum ada"
            )

        wb = openpyxl.Workbook()
        ws = wb.active
        CURRENT_SHEET_NAME = make_unique_sheet_name(wb, requested_name)
        ws.title = CURRENT_SHEET_NAME
        COMPARISON_SHEET_NAMES = []
        existing_tags = set()
        format_sheet(ws)
        wb.save(TAGS_FILE)

    wb.close()
    return len(existing_tags), len(seen_tags)


def sync_excel():
    """Sinkronkan hasil sesi saat ini tanpa mengubah sheet-sheet lama."""
    wb = openpyxl.load_workbook(TAGS_FILE)
    ws = wb[CURRENT_SHEET_NAME]

    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    format_sheet(ws)

    zebra_fill = PatternFill("solid", fgColor="D9E1F2")

    # Urutkan tag_id secara ascending
    sorted_tag_ids = sorted(seen_tags.keys())

    for idx, tag_id in enumerate(sorted_tag_ids, 1):
        info = seen_tags[tag_id]
        row_idx = idx + 1
        ws.append([idx, tag_id, info["device_id"]])

        if idx % 2 == 0:
            for col in range(1, len(HEADERS) + 1):
                ws.cell(row=row_idx, column=col).fill = zebra_fill

    wb.save(TAGS_FILE)
    wb.close()


def save_or_update_tag(tag_id, parsed):
    """Simpan tag yang tidak ada di seluruh sheet lama atau sesi saat ini."""
    device_id = parsed.get("device_id")

    if NO_SAVE:
        if tag_id not in seen_tags:
            seen_tags[tag_id] = {"device_id": device_id}
            session_new_tags.add(tag_id)
            return "NEW"
        return "SKIPPED"

    if tag_id in existing_tags:
        return "EXISTS_IN_WORKBOOK"

    if tag_id not in seen_tags:
        seen_tags[tag_id] = {"device_id": device_id}
        session_new_tags.add(tag_id)
        sync_excel()
        return "NEW"

    return "SKIPPED"


def now():
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S.%f")[:-3]


def log(message):
    print(f"[{now()}] {message}", flush=True)


def format_hex(data):
    return data.hex(" ").upper()


def extract_frames(buffer):
    """Extract frames menggunakan length byte atau 0D0A end marker."""
    frames = []

    while True:
        start = buffer.find(b"\xbb")
        if start == -1:
            return frames, b""

        if start > 0:
            buffer = buffer[start:]

        if len(buffer) < 6:
            return frames, buffer

        data_len = buffer[2]
        frame_len = data_len + 6

        # Jika buffer cukup panjang untuk frame berbasis length byte
        if len(buffer) >= frame_len and buffer[frame_len - 2] == 0x0D and buffer[frame_len - 1] == 0x0A:
            frames.append(buffer[:frame_len])
            buffer = buffer[frame_len:]
            continue

        # Fallback: cari 0D 0A terdekat jika length byte tidak cocok
        end = buffer.find(b"\x0d\x0a")
        if end != -1:
            frames.append(buffer[:end + 2])
            buffer = buffer[end + 2:]
        else:
            return frames, buffer


def parse_frame(frame):
    if len(frame) < 4:
        return {
            "protocol": "unknown",
            "raw": format_hex(frame),
        }

    command = frame[1]

    if command == 0xF1 and len(frame) >= 17:
        tag_id = frame[4:9].hex().upper()
        device_id = frame[9:16].hex().upper()
        power = frame[3]

        return {
            "protocol": "2.4G",
            "command": "F1",
            "id": tag_id,
            "epc": tag_id,
            "tid": None,
            "device_id": device_id,
            "power": power,
            "raw": format_hex(frame),
        }

    payload = frame[3:-3] if len(frame) >= 6 else b""
    command_name = {
        0x16: "START_SINGLE_EPC",
        0x17: "START_MULTI_EPC",
        0x18: "STOP_MULTI_EPC",
        0x79: "MULTI_TAGS_USER",
        0x96: "START_SINGLE_EPC_RESPOND",
        0x97: "START_MULTI_EPC_RESPOND",
        0x98: "STOP_MULTI_EPC_RESPOND",
        0xF1: "GET_MULTI_ID",
    }.get(command, f"0x{command:02X}")

    return {
        "protocol": "UHF/unknown",
        "command": command_name,
        "payload": format_hex(payload),
        "raw": format_hex(frame),
    }

def print_tag(parsed):
    if parsed.get("protocol") == "2.4G":
        tag_id = parsed.get("epc")

        # Filter Tag ID Pengecualian
        if tag_id in EXCLUDED_TAGS:
            return

        # Tag ID valid hanya berisi angka. Abaikan hasil scan yang mengandung huruf.
        if not tag_id or not tag_id.isdigit():
            return

        # Filter Range Tag ID
        if TAG_MIN and tag_id < TAG_MIN:
            return
        if TAG_MAX and tag_id > TAG_MAX:
            return

        action = save_or_update_tag(tag_id, parsed)

        if action == "NEW":
            label = "TEST TAG" if NO_SAVE else "NEW TAG"
            log(
                f"[{label} #{len(session_new_tags)}] "
                f"EPC={tag_id} "
                f"DEVICE={parsed.get('device_id')}"
            )
        return

    log(
        "FRAME "
        f"COMMAND={parsed.get('command')} "
        f"PAYLOAD={parsed.get('payload')} "
        f"RAW={parsed.get('raw')}"
    )


def send_command(sock, command, label):
    try:
        sock.sendall(command)
        log(f"TX {label}: {format_hex(command)}")
        return True
    except OSError as exc:
        log(f"TX {label} failed: {exc}")
        return False


def run_stop_only(args):
    """Kirim perintah STOP tanpa membuat/mengubah sheet Excel."""
    if args.mode == "client":
        log(f"Connecting to RFID device {args.reader_ip}:{args.port} to stop scan")
        with socket.create_connection(
            (args.reader_ip, args.port), timeout=args.connect_timeout
        ) as sock:
            sock.settimeout(1)
            if not send_command(sock, STOP_COMMAND, "stop-only"):
                raise OSError("Gagal mengirim perintah STOP")
        log("Perintah STOP terkirim; receiver seharusnya berhenti membaca tag")
        return

    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as server:
        server.setsockopt(socket.SOL_SOCKET, socket.SO_REUSEADDR, 1)
        server.bind((args.host, args.port))
        server.listen(1)
        server.settimeout(args.connect_timeout)
        log(f"Waiting for RFID device on {args.host}:{args.port} to stop scan")
        sock, address = server.accept()
        with sock:
            log(f"Connected ({address[0]}:{address[1]})")
            if not send_command(sock, STOP_COMMAND, "stop-only"):
                raise OSError("Gagal mengirim perintah STOP")
        log("Perintah STOP terkirim; receiver seharusnya berhenti membaca tag")


def receive_frames(sock):
    buffer = b""
    sock.settimeout(1)

    send_command(sock, START_COMMAND, "auto-start")

    while True:
        try:
            data = sock.recv(4096)
        except socket.timeout:
            continue

        if not data:
            log("Disconnected")
            return

        log(f"RX {len(data)} bytes: {format_hex(data)}")
        buffer += data
        frames, buffer = extract_frames(buffer)

        for frame in frames:
            parsed = parse_frame(frame)
            print_tag(parsed)


def run_server(args):
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as server:
        server.setsockopt(socket.SOL_SOCKET, socket.SO_REUSEADDR, 1)
        server.bind((args.host, args.port))
        server.listen(1)
        server.settimeout(1)

        while True:
            log("Status: Disconnected")
            log(f"Listening on {args.host}:{args.port}; waiting for RFID device")

            while True:
                try:
                    sock, address = server.accept()
                    break
                except socket.timeout:
                    continue

            with sock:
                log(f"Status: Connected ({address[0]}:{address[1]})")
                try:
                    receive_frames(sock)
                finally:
                    send_command(sock, STOP_COMMAND, "auto-stop")
                    log("Connection closed")


def run_client(args):
    log(f"Status: Disconnected")
    log(f"Connecting to RFID device {args.reader_ip}:{args.port}")

    with socket.create_connection((args.reader_ip, args.port), timeout=args.connect_timeout) as sock:
        log(f"Status: Connected ({args.reader_ip}:{args.port})")
        try:
            receive_frames(sock)
        finally:
            send_command(sock, STOP_COMMAND, "auto-stop")
            log("Connection closed")


def parse_args():
    global TAG_MIN, TAG_MAX, NO_SAVE
    parser = argparse.ArgumentParser()
    parser.add_argument("--mode", choices=["server", "client"], default=os.getenv("RFID_MODE", "server"))
    parser.add_argument("--host", default=os.getenv("RFID_LISTEN_HOST", "0.0.0.0"))
    # Baca RFID_READER_1_HOST (format baru) atau RFID_READER_HOST (format lama)
    default_ip = os.getenv("RFID_READER_1_HOST") or os.getenv("RFID_READER_HOST", "172.15.0.80")
    parser.add_argument("--reader-ip", default=default_ip)
    parser.add_argument("--port", type=int, default=int(os.getenv("RFID_PORT", "6000")))
    parser.add_argument("--connect-timeout", type=float, default=float(os.getenv("RFID_CONNECT_TIMEOUT", "5")))
    parser.add_argument(
        "--stop-only",
        action="store_true",
        help="Kirim perintah STOP ke receiver tanpa membuat atau mengubah sheet Excel",
    )
    parser.add_argument(
        "--remove-rssi-only",
        action="store_true",
        help="Hapus kolom RSSI dari seluruh sheet lalu keluar tanpa scan",
    )
    parser.add_argument("--tag-min", default=os.getenv("RFID_TAG_MIN", ""), help="Minimum Tag ID")
    parser.add_argument("--tag-max", default=os.getenv("RFID_TAG_MAX", ""), help="Maximum Tag ID")
    sheet_group = parser.add_mutually_exclusive_group()
    sheet_group.add_argument(
        "--sheet-name",
        default=os.getenv("RFID_SHEET_NAME"),
        help="Nama sheet untuk sesi scan baru (default: Scan YYYYMMDD HHMMSS)",
    )
    sheet_group.add_argument(
        "--continue-sheet",
        help="Lanjutkan scan ke sheet yang sudah ada tanpa menghapus data lamanya",
    )
    sheet_group.add_argument(
        "--no-save",
        action="store_true",
        help="Tes reader tanpa membaca, membuat, atau mengubah sheet Excel",
    )
    args = parser.parse_args()

    TAG_MIN = args.tag_min
    TAG_MAX = args.tag_max
    NO_SAVE = args.no_save

    if TAG_MIN or TAG_MAX:
        log(f"Range Tag ID filter aktif: {TAG_MIN or 'MIN'} s/d {TAG_MAX or 'MAX'}")

    return args


def main():
    args = parse_args()

    if args.stop_only:
        try:
            run_stop_only(args)
        except (OSError, socket.timeout) as exc:
            log(f"Gagal menghentikan receiver: {exc}")
            sys.exit(1)
        return

    if args.remove_rssi_only:
        try:
            changed_sheets = remove_rssi_from_existing_workbook()
        except (
            OSError,
            KeyError,
            openpyxl.utils.exceptions.InvalidFileException,
        ) as exc:
            log(f"Gagal mengubah file Excel {os.path.abspath(TAGS_FILE)}: {exc}")
            sys.exit(1)

        if changed_sheets:
            log("Kolom RSSI dihapus dari: " + ", ".join(changed_sheets))
            log(f"Workbook tersimpan: {os.path.abspath(TAGS_FILE)}")
        else:
            log("Tidak ada kolom RSSI yang perlu dihapus")
        return

    if args.no_save:
        seen_tags.clear()
        existing_tags.clear()
        session_new_tags.clear()
        log("Mode test reader: hasil scan hanya tampil di terminal dan tidak disimpan")
    else:
        try:
            existing_count, target_count = initialize_scan_sheet(
                requested_name=args.sheet_name,
                continue_sheet=args.continue_sheet,
            )
        except (OSError, ValueError, KeyError, openpyxl.utils.exceptions.InvalidFileException) as exc:
            log(f"Gagal membaca file Excel {os.path.abspath(TAGS_FILE)}: {exc}")
            sys.exit(1)

        if COMPARISON_SHEET_NAMES:
            log(
                f"Sheet pembanding: semua {len(COMPARISON_SHEET_NAMES)} sheet lama "
                f"({existing_count} unique tags): {', '.join(COMPARISON_SHEET_NAMES)}"
            )
        else:
            log("Sheet pembanding: tidak ada (workbook baru)")
        if args.continue_sheet:
            log(f"Melanjutkan sheet: {CURRENT_SHEET_NAME} ({target_count} tag tersimpan)")
        else:
            log(f"Sheet scan baru: {CURRENT_SHEET_NAME}")

    try:
        if args.mode == "server":
            run_server(args)
        else:
            run_client(args)
    except KeyboardInterrupt:
        log("Interrupted by Ctrl+C")
    except OSError as exc:
        log(f"Socket error: {exc}")
        sys.exit(1)
    finally:
        log("Status: Disconnected")
        if args.no_save:
            log(f"Total Tag ID unik terdeteksi (tidak disimpan): {len(session_new_tags)}")
            for i, t in enumerate(sorted(session_new_tags), 1):
                log(f"  {i}. {t}")
        else:
            log(f"Total tag baru di sheet {CURRENT_SHEET_NAME}: {len(session_new_tags)}")
            if session_new_tags:
                log(f"Tags saved to: {os.path.abspath(TAGS_FILE)}")
                for i, t in enumerate(sorted(session_new_tags), 1):
                    log(f"  {i}. {t}")


if __name__ == "__main__":
    def handle_termination(signum, frame):
        del signum, frame
        raise KeyboardInterrupt

    signal.signal(signal.SIGTERM, handle_termination)
    if hasattr(signal, "SIGHUP"):
        signal.signal(signal.SIGHUP, handle_termination)
    main()
